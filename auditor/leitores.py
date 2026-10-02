"""Leitura dos arquivos anexados (PDF, OFX, DOC/DOCX, XLS/XLSX, CSV) e extração de valores.

A extração é heurística: relatórios de cada sistema têm layouts diferentes. Por isso a
interface sempre mostra a tabela mensal consolidada para o analista conferir e corrigir.
"""

from __future__ import annotations

import io
import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime

import pandas as pd

from .modelos import Documento, Lancamento, TipoDocumento

# ---------------------------------------------------------------------------
# Utilidades de texto, números e datas
# ---------------------------------------------------------------------------

MESES = {
    "jan": 1, "fev": 2, "mar": 3, "abr": 4, "mai": 5, "jun": 6,
    "jul": 7, "ago": 8, "set": 9, "out": 10, "nov": 11, "dez": 12,
}

RE_VALOR = re.compile(
    r"(?<![\d/])(?:R\$\s*)?(\(?-?\s?\d{1,3}(?:\.\d{3})*,\d{2}\)?|\(?-?\s?\d+,\d{2}\)?)(\s?[CD](?![A-Za-z]))?"
)
RE_DATA_COMPLETA = re.compile(r"\b(\d{2})[/.-](\d{2})[/.-](\d{4}|\d{2})\b")
RE_DATA_CURTA = re.compile(r"^\s*(\d{2})/(\d{2})(?!/\d)")
RE_COMPETENCIA = re.compile(
    r"(?:compet[eê]ncia|per[ií]odo de apura[cç][aã]o|m[eê]s/ano|refer[eê]ncia|per[ií]odo)\s*[:\-]?\s*(\d{2})/(\d{4})",
    re.IGNORECASE,
)
RE_MES_EXTENSO = re.compile(
    r"\b(jan|fev|mar|abr|mai|jun|jul|ago|set|out|nov|dez)[a-zç]*\.?\s*(?:/|de|-)?\s*(\d{4}|\d{2})\b",
    re.IGNORECASE,
)


def normalizar(texto: str) -> str:
    sem_acento = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return sem_acento.lower().strip()


def parse_valor(bruto) -> float | None:
    """Converte '1.234,56', '(1.234,56)', '1.234,56 D', '-1234.56', 1234.5 para float."""
    if bruto is None:
        return None
    if isinstance(bruto, (int, float)):
        if pd.isna(bruto):
            return None
        return float(bruto)
    s = str(bruto).strip()
    if not s:
        return None
    negativo = False
    if s.startswith("(") and s.endswith(")"):
        negativo, s = True, s[1:-1]
    s = s.replace("R$", "").replace(" ", "").strip()
    if s.upper().endswith("D"):
        negativo, s = True, s[:-1]
    elif s.upper().endswith("C"):
        s = s[:-1]
    if s.endswith("-"):
        negativo, s = True, s[:-1]
    if s.startswith("-"):
        negativo, s = True, s[1:]
    if not re.fullmatch(r"[\d.,]+", s):
        return None
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif s.count(".") > 1:
        s = s.replace(".", "")
    try:
        v = float(s)
    except ValueError:
        return None
    return -v if negativo else v


def parse_data(bruto, ano_padrao: int | None = None) -> date | None:
    if bruto is None or (isinstance(bruto, float) and pd.isna(bruto)):
        return None
    if isinstance(bruto, datetime):
        return bruto.date()
    if isinstance(bruto, date):
        return bruto
    if isinstance(bruto, pd.Timestamp):
        return bruto.date()
    s = str(bruto).strip()
    m = RE_DATA_COMPLETA.search(s)
    if m:
        d, mes, a = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if a < 100:
            a += 2000
        try:
            return date(a, mes, d)
        except ValueError:
            return None
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return None
    m = RE_DATA_CURTA.match(s)
    if m and ano_padrao:
        try:
            return date(ano_padrao, int(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
    return None


def parse_competencia(bruto) -> str | None:
    """Reconhece '01/2025', 'jan/25', 'Janeiro de 2025', '2025-01', datas completas."""
    if bruto is None or (isinstance(bruto, float) and pd.isna(bruto)):
        return None
    if isinstance(bruto, (datetime, date, pd.Timestamp)):
        return f"{bruto.year}-{bruto.month:02d}"
    s = normalizar(bruto)
    m = re.fullmatch(r"(\d{1,2})[/.-](\d{4})", s)
    if m and 1 <= int(m.group(1)) <= 12:
        return f"{m.group(2)}-{int(m.group(1)):02d}"
    m = re.fullmatch(r"(\d{4})[/.-](\d{1,2})", s)
    if m and 1 <= int(m.group(2)) <= 12:
        return f"{m.group(1)}-{int(m.group(2)):02d}"
    m = RE_MES_EXTENSO.fullmatch(s) or RE_MES_EXTENSO.match(s)
    if m:
        a = int(m.group(2))
        a = a + 2000 if a < 100 else a
        return f"{a}-{MESES[m.group(1)[:3].lower()]:02d}"
    d = parse_data(s)
    if d:
        return f"{d.year}-{d.month:02d}"
    return None


def competencia(d: date) -> str:
    return f"{d.year}-{d.month:02d}"


# ---------------------------------------------------------------------------
# Leitura bruta por formato
# ---------------------------------------------------------------------------


def _ler_pdf(conteudo: bytes) -> tuple[str, list[pd.DataFrame], list[str]]:
    import pdfplumber

    textos, tabelas, avisos = [], [], []
    with pdfplumber.open(io.BytesIO(conteudo)) as pdf:
        for pagina in pdf.pages:
            textos.append(pagina.extract_text() or "")
            for t in pagina.extract_tables() or []:
                if t and len(t) > 1:
                    cab = [str(c or f"col{i}") for i, c in enumerate(t[0])]
                    tabelas.append(pd.DataFrame(t[1:], columns=_unicos(cab)))
    texto = "\n".join(textos)
    if len(texto.strip()) < 20:
        avisos.append(
            "O PDF parece ser uma imagem escaneada (sem texto). Use a análise com IA ou informe os valores manualmente."
        )
    return texto, tabelas, avisos


def _ler_docx(conteudo: bytes) -> tuple[str, list[pd.DataFrame]]:
    import docx

    doc = docx.Document(io.BytesIO(conteudo))
    linhas = [p.text for p in doc.paragraphs]
    tabelas = []
    for t in doc.tables:
        dados = [[c.text for c in r.cells] for r in t.rows]
        if len(dados) > 1:
            tabelas.append(pd.DataFrame(dados[1:], columns=_unicos(dados[0])))
            linhas.extend(" ".join(r) for r in dados)
    return "\n".join(linhas), tabelas


def _ler_planilha(conteudo: bytes, extensao: str) -> list[pd.DataFrame]:
    engine = "xlrd" if extensao == "xls" else None
    planilhas = pd.read_excel(io.BytesIO(conteudo), sheet_name=None, header=None, engine=engine)
    return [_promover_cabecalho(df) for df in planilhas.values() if not df.dropna(how="all").empty]


def _ler_csv(conteudo: bytes) -> list[pd.DataFrame]:
    for enc in ("utf-8-sig", "latin-1"):
        try:
            texto = conteudo.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    primeira = texto.splitlines()[0] if texto else ""
    sep = ";" if primeira.count(";") >= primeira.count(",") else ","
    df = pd.read_csv(io.StringIO(texto), sep=sep, header=None, dtype=str, on_bad_lines="skip")
    return [_promover_cabecalho(df)]


def _unicos(colunas) -> list[str]:
    vistos: dict[str, int] = {}
    saida = []
    for c in colunas:
        c = str(c).strip() if c is not None and str(c).strip() else "col"
        if c in vistos:
            vistos[c] += 1
            c = f"{c}_{vistos[c]}"
        else:
            vistos[c] = 0
        saida.append(c)
    return saida


def _promover_cabecalho(df: pd.DataFrame) -> pd.DataFrame:
    """Usa como cabeçalho a primeira linha com maioria de células de texto (relatórios têm títulos antes)."""
    df = df.dropna(how="all").dropna(axis=1, how="all").reset_index(drop=True)
    for i in range(min(15, len(df))):
        linha = df.iloc[i]
        preenchidas = linha.dropna()
        textos = [v for v in preenchidas if isinstance(v, str) and parse_valor(v) is None and parse_data(v) is None]
        if len(preenchidas) >= 2 and len(textos) >= max(2, int(len(df.columns) * 0.5)):
            novo = df.iloc[i + 1 :].copy()
            novo.columns = _unicos(linha.tolist())
            return novo.reset_index(drop=True)
    df.columns = [f"col{i}" for i in range(len(df.columns))]
    return df


def _ler_ofx(conteudo: bytes, nome: str) -> list[Lancamento]:
    texto = conteudo.decode("latin-1", errors="ignore")
    lancs = []
    for bloco in re.findall(r"<STMTTRN>(.*?)(?:</STMTTRN>|(?=<STMTTRN>)|</BANKTRANLIST>)", texto, re.S | re.I):
        def campo(tag):
            m = re.search(rf"<{tag}>([^<\r\n]*)", bloco, re.I)
            return m.group(1).strip() if m else ""

        valor = parse_valor(campo("TRNAMT").replace(",", "."))
        dt = campo("DTPOSTED")
        if valor is None or len(dt) < 8:
            continue
        try:
            d = date(int(dt[:4]), int(dt[4:6]), int(dt[6:8]))
        except ValueError:
            continue
        desc = " ".join(x for x in (campo("NAME"), campo("MEMO")) if x) or campo("TRNTYPE")
        lancs.append(Lancamento(d, desc, valor, nome))
    return lancs


# ---------------------------------------------------------------------------
# Detecção do tipo de documento
# ---------------------------------------------------------------------------

PALAVRAS_TIPO = [
    (TipoDocumento.RENDIMENTOS, ["informe de rendimentos", "rendimentos financeiros", "comprovante de rendimentos"]),
    (TipoDocumento.EXTRATO_APLICACAO, ["aplicacao", "cdb", "fundo de investimento", "lci", "lca", "poupanca", "resgate"]),
    (TipoDocumento.FOLHA, ["folha de pagamento", "resumo da folha", "proventos", "inss segurado", "fgts", "pro-labore", "pro labore"]),
    (TipoDocumento.FISCAL, ["pgdas", "totais icms por natureza", "registro de entradas", "registro de saidas", "extrato do simples", "dctf", "apuracao", "sped", "livro de registro", "efd", "darf"]),
    (TipoDocumento.DEMONSTRATIVO, ["balanco patrimonial", "balancete", "demonstracao do resultado", "dre", "ativo circulante"]),
    (TipoDocumento.COMPRAS, ["compras", "entradas", "fornecedor", "notas de entrada"]),
    (TipoDocumento.FATURAMENTO, ["faturamento", "vendas", "saidas", "notas fiscais emitidas", "receita"]),
    (TipoDocumento.EXTRATO_CC, ["extrato", "conta corrente", "saldo anterior", "saldo do dia", "agencia"]),
]


def detectar_tipo(nome: str, texto: str = "") -> TipoDocumento:
    if nome.lower().endswith(".ofx"):
        alvo = normalizar(texto[:3000])
        return TipoDocumento.EXTRATO_APLICACAO if "invstmt" in alvo else TipoDocumento.EXTRATO_CC
    alvo = normalizar(nome) + " " + normalizar(texto[:4000])
    melhor, pontos = TipoDocumento.OUTRO, 0
    for tipo, palavras in PALAVRAS_TIPO:
        p = sum(1 for w in palavras if w in alvo)  # presença, não frequência: linhas repetidas não distorcem
        if normalizar(nome).find(palavras[0].split()[0]) >= 0:
            p += 3
        if p > pontos:
            melhor, pontos = tipo, p
    return melhor


# ---------------------------------------------------------------------------
# Extração de lançamentos (extratos)
# ---------------------------------------------------------------------------

PALAVRAS_DEBITO = [
    "pagamento", "pagto", "pgto", "tarifa", "debito", "deb ", "saque", "compra", "pix enviado",
    "ted enviada", "doc enviado", "transf enviada", "aplicacao", "boleto pago", "tributo", "darf", "das ",
    "folha", "salario", "iof", "juros", "encargo",
]


def _sinal_por_descricao(desc: str) -> int:
    d = normalizar(desc)
    return -1 if any(p in d for p in PALAVRAS_DEBITO) else 1


def _ano_do_texto(texto: str, ano_padrao: int) -> int:
    anos = re.findall(r"\b\d{2}/\d{2}/(20\d{2})\b", texto)
    if anos:
        return int(max(set(anos), key=anos.count))
    return ano_padrao


def lancamentos_de_texto(texto: str, nome: str, ano_padrao: int) -> list[Lancamento]:
    ano = _ano_do_texto(texto, ano_padrao)
    lancs = []
    for linha in texto.splitlines():
        if "saldo" in normalizar(linha):
            continue
        d = None
        m = RE_DATA_COMPLETA.match(linha.strip())
        if m:
            d = parse_data(m.group(0))
            resto = linha.strip()[m.end():]
        else:
            m = RE_DATA_CURTA.match(linha)
            if m:
                d = parse_data(m.group(0), ano)
                resto = linha[m.end():]
        if not d:
            continue
        valores = list(RE_VALOR.finditer(resto))
        if not valores:
            continue
        v = valores[0]
        bruto = v.group(1) + (v.group(2) or "")
        valor = parse_valor(bruto)
        if valor is None or valor == 0:
            continue
        desc = resto[: v.start()].strip()
        explicito = v.group(2) or "-" in v.group(1) or "(" in v.group(1)
        if not explicito:
            valor = abs(valor) * _sinal_por_descricao(desc)
        lancs.append(Lancamento(d, desc, valor, nome))
    return lancs


def _achar_coluna(colunas, palavras) -> str | None:
    for c in colunas:
        n = normalizar(c)
        if any(p in n for p in palavras):
            return c
    return None


def lancamentos_de_tabela(df: pd.DataFrame, nome: str, ano_padrao: int) -> list[Lancamento]:
    cols = list(df.columns)
    c_data = _achar_coluna(cols, ["data", "dt ", "date"]) or _coluna_de_datas(df, ano_padrao)
    if not c_data:
        return []
    c_desc = _achar_coluna(cols, ["historico", "descricao", "lancamento", "memo", "detalhe"])
    c_cred = _achar_coluna(cols, ["credito", "entrada"])
    c_deb = _achar_coluna(cols, ["debito", "saida"])
    c_valor = _achar_coluna(cols, ["valor", "montante", "quantia"])
    c_tipo = _achar_coluna(cols, ["tipo", "d/c", "c/d", "natureza"])
    lancs = []
    for _, r in df.iterrows():
        d = parse_data(r[c_data], ano_padrao)
        if not d:
            continue
        desc = str(r[c_desc]) if c_desc else ""
        if "saldo" in normalizar(desc):
            continue
        valor = None
        if c_cred or c_deb:
            cr = parse_valor(r[c_cred]) if c_cred else None
            db = parse_valor(r[c_deb]) if c_deb else None
            if cr:
                valor = abs(cr)
            elif db:
                valor = -abs(db)
        elif c_valor:
            valor = parse_valor(r[c_valor])
            if valor is not None and c_tipo and normalizar(r[c_tipo]).startswith("d"):
                valor = -abs(valor)
        if valor:
            lancs.append(Lancamento(d, desc, valor, nome))
    return lancs


def _coluna_de_datas(df: pd.DataFrame, ano_padrao: int) -> str | None:
    for c in df.columns:
        amostra = df[c].dropna().head(20)
        if len(amostra) and sum(parse_data(v, ano_padrao) is not None for v in amostra) >= len(amostra) * 0.6:
            return c
    return None


# ---------------------------------------------------------------------------
# Extração de séries mensais (faturamento, compras, folha, rendimentos, fiscal)
# ---------------------------------------------------------------------------


def serie_de_tabela(df: pd.DataFrame, ano_padrao: int, palavras_valor: list[str] | None = None) -> dict[str, float]:
    cols = list(df.columns)
    c_comp = _achar_coluna(cols, ["competencia", "mes", "periodo", "referencia"])
    c_data = c_comp or _achar_coluna(cols, ["data", "emissao", "dt "]) or _coluna_de_datas(df, ano_padrao)
    if not c_data:
        return {}
    c_valor = None
    if palavras_valor:
        c_valor = _achar_coluna(cols, palavras_valor)
    c_valor = c_valor or _achar_coluna(cols, ["valor total", "total", "valor contabil", "valor", "vlr", "bruto", "receita"])
    if not c_valor:
        numericas = [c for c in cols if c != c_data and df[c].map(parse_valor).notna().mean() > 0.6]
        if not numericas:
            return {}
        c_valor = max(numericas, key=lambda c: df[c].map(parse_valor).abs().sum())
    serie: dict[str, float] = defaultdict(float)
    for _, r in df.iterrows():
        comp = parse_competencia(r[c_data]) if c_comp else None
        if not comp:
            d = parse_data(r[c_data], ano_padrao)
            comp = competencia(d) if d else None
        v = parse_valor(r[c_valor])
        linha = " ".join(normalizar(x) for x in r.astype(str).tolist())
        if comp and v is not None and "total" not in linha:
            serie[comp] += v
    return dict(serie)


def serie_de_texto(texto: str, ano_padrao: int) -> dict[str, float]:
    """Soma linhas com data + valor por mês; se não houver, usa 'Total' + competência do documento."""
    serie: dict[str, float] = defaultdict(float)
    for linha in texto.splitlines():
        n = normalizar(linha)
        if "total" in n or "saldo" in n:
            continue
        m = RE_DATA_COMPLETA.search(linha)
        if not m:
            continue
        d = parse_data(m.group(0))
        valores = [parse_valor(v.group(1)) for v in RE_VALOR.finditer(linha[m.end():])]
        valores = [v for v in valores if v is not None]
        if d and valores:
            serie[competencia(d)] += max(valores, key=abs)
    if serie:
        return dict(serie)
    comp = _competencia_do_texto(texto)
    total = valor_por_palavras(texto, ["valor total", "total geral", "total"])
    if comp and total is not None:
        return {comp: total}
    return {}


def _competencia_do_texto(texto: str) -> str | None:
    m = RE_COMPETENCIA.search(texto)
    if m:
        return f"{m.group(2)}-{m.group(1)}"
    m = RE_MES_EXTENSO.search(normalizar(texto[:2000]))
    if m:
        return parse_competencia(m.group(0))
    return None


def valor_por_palavras(texto: str, palavras: list[str]) -> float | None:
    """Último valor monetário da primeira linha que contém alguma das palavras (em ordem de prioridade)."""
    linhas = [(normalizar(l), l) for l in texto.splitlines()]
    for p in palavras:
        for n, original in linhas:
            if p in n:
                valores = [parse_valor(v.group(1) + (v.group(2) or "")) for v in RE_VALOR.finditer(original)]
                valores = [v for v in valores if v is not None]
                if valores:
                    return valores[-1]
    return None


def _blocos_por_competencia(texto: str) -> list[tuple[str, str]]:
    """Divide relatórios com vários meses (ex.: folha anual) em blocos por competência."""
    marcas = list(RE_COMPETENCIA.finditer(texto))
    if not marcas:
        comp = _competencia_do_texto(texto)
        return [(comp, texto)] if comp else []
    blocos = []
    for i, m in enumerate(marcas):
        fim = marcas[i + 1].start() if i + 1 < len(marcas) else len(texto)
        blocos.append((f"{m.group(2)}-{m.group(1)}", texto[m.start():fim]))
    return blocos


CONTAS_DEMONSTRATIVO = {
    "caixa": ["caixa geral", "caixa "],
    "bancos": ["bancos conta movimento", "bancos c/ movimento", "bancos"],
    "aplicacoes": ["aplicacoes financeiras", "aplicacao financeira"],
    "estoques": ["estoques", "estoque de mercadorias"],
    "total_ativo": ["total do ativo", "ativo total", "total ativo"],
    "total_passivo": ["total do passivo", "passivo total", "total passivo"],
    "patrimonio_liquido": ["total do patrimonio liquido", "patrimonio liquido"],
    "capital_social": ["capital social"],
    "emprestimos_socios": ["emprestimos de socios", "mutuo", "conta corrente de socios", "adiantamento de socios"],
    "receita_bruta": ["receita bruta", "receita operacional bruta", "receita de vendas"],
    "cmv": ["custo das mercadorias", "custo dos produtos", "custo dos servicos", "cmv", "cpv", "csp"],
    "despesas_pessoal": ["despesas com pessoal", "salarios e ordenados", "salarios"],
    "pro_labore": ["pro-labore", "pro labore"],
    "receitas_financeiras": ["receitas financeiras", "rendimentos de aplicacoes"],
    "lucro_liquido": ["lucro liquido do exercicio", "lucro liquido", "resultado liquido", "lucro do periodo"],
    "prejuizo": ["prejuizo do exercicio", "prejuizo liquido"],
    "distribuicao_lucros": ["lucros distribuidos", "distribuicao de lucros", "dividendos"],
    "impostos_resultado": ["provisao para irpj", "irpj", "contribuicao social"],
}


def extrair_contas(texto: str) -> dict[str, float]:
    contas = {}
    for chave, palavras in CONTAS_DEMONSTRATIVO.items():
        v = valor_por_palavras(texto, palavras)
        if v is not None:
            contas[chave] = v
    return contas


def _series_folha(texto: str, tabelas: list[pd.DataFrame], ano: int) -> dict[str, dict[str, float]]:
    series: dict[str, dict[str, float]] = defaultdict(dict)
    for comp, bloco in _blocos_por_competencia(texto):
        sal = valor_por_palavras(bloco, ["total de proventos", "total proventos", "total bruto", "salario contribuicao", "base inss", "salarios"])
        pro = valor_por_palavras(bloco, ["pro-labore", "pro labore", "prolabore"])
        inss = valor_por_palavras(bloco, ["inss empresa", "inss patronal", "cpp", "contribuicao patronal"])
        fgts = valor_por_palavras(bloco, ["fgts"])
        if sal is not None:
            if pro is not None and pro < sal:
                sal -= pro  # total de proventos geralmente já inclui o pró-labore
            series["folha_salarios"][comp] = series["folha_salarios"].get(comp, 0) + sal
        if pro is not None:
            series["pro_labore"][comp] = series["pro_labore"].get(comp, 0) + pro
        enc = (inss or 0) + (fgts or 0)
        if enc:
            series["encargos_folha"][comp] = series["encargos_folha"].get(comp, 0) + enc
    if not series:
        for df in tabelas:
            s = serie_de_tabela(df, ano, ["proventos", "bruto", "salario"])
            if s:
                series["folha_salarios"] = s
                break
    return dict(series)


def _series_fiscal(texto: str, tabelas: list[pd.DataFrame], ano: int) -> dict[str, dict[str, float]]:
    from .layouts import pgdas_d

    especifico = pgdas_d(texto)
    if especifico:
        return especifico
    series: dict[str, dict[str, float]] = defaultdict(dict)
    for comp, bloco in _blocos_por_competencia(texto):
        rec = valor_por_palavras(bloco, ["receita bruta do pa", "receita bruta auferida", "receita bruta", "total de receitas", "faturamento"])
        imp = valor_por_palavras(bloco, ["valor total do debito", "total do debito", "total a recolher", "valor do das", "total devido", "valor a pagar"])
        if rec is not None:
            series["faturamento_declarado"][comp] = rec
        if imp is not None:
            series["imposto_declarado"][comp] = imp
    if not series:
        for df in tabelas:
            s = serie_de_tabela(df, ano, ["receita", "faturamento"])
            if s:
                series["faturamento_declarado"] = s
                break
    return dict(series)


def _series_rendimentos(doc: Documento, ano: int) -> dict[str, dict[str, float]]:
    series: dict[str, dict[str, float]] = defaultdict(dict)
    for l in doc.lancamentos:
        n = normalizar(l.descricao)
        comp = competencia(l.data)
        if "ir" in n.split() or "irrf" in n or "imposto de renda" in n:
            series["irrf_aplicacao"][comp] = series["irrf_aplicacao"].get(comp, 0) + abs(l.valor)
        elif any(p in n for p in ("rendimento", "juros", "remuneracao", "rend ")):
            series["rendimentos_aplicacao"][comp] = series["rendimentos_aplicacao"].get(comp, 0) + abs(l.valor)
    if series:
        return dict(series)
    rend = valor_por_palavras(doc.texto, ["rendimento bruto", "rendimentos", "rendimento"])
    irrf = valor_por_palavras(doc.texto, ["imposto de renda retido", "irrf", "ir retido"])
    if rend is not None:
        # Informe anual: distribui igualmente no ano quando não há detalhamento mensal.
        comp = _competencia_do_texto(doc.texto)
        alvo = [comp] if comp else [f"{ano}-{m:02d}" for m in range(1, 13)]
        for c in alvo:
            series["rendimentos_aplicacao"][c] = abs(rend) / len(alvo)
            if irrf is not None:
                series["irrf_aplicacao"][c] = abs(irrf) / len(alvo)
        if not comp:
            doc.avisos.append("Informe sem detalhe mensal: rendimentos distribuídos igualmente nos 12 meses.")
    return dict(series)


# ---------------------------------------------------------------------------
# Ponto de entrada
# ---------------------------------------------------------------------------

EXTENSOES = ("pdf", "ofx", "doc", "docx", "xls", "xlsx", "csv")


def ler_documento(nome: str, conteudo: bytes, tipo: TipoDocumento | None = None, ano_padrao: int | None = None) -> Documento:
    ano = ano_padrao or date.today().year
    ext = nome.lower().rsplit(".", 1)[-1] if "." in nome else ""
    texto, tabelas, avisos, lancs = "", [], [], []

    try:
        if ext == "pdf":
            texto, tabelas, avisos = _ler_pdf(conteudo)
        elif ext == "docx":
            texto, tabelas = _ler_docx(conteudo)
        elif ext == "doc":
            avisos.append("Arquivos .DOC antigos não são lidos diretamente. Abra no Word e salve como .DOCX ou PDF.")
        elif ext in ("xls", "xlsx"):
            tabelas = _ler_planilha(conteudo, ext)
        elif ext == "csv":
            tabelas = _ler_csv(conteudo)
        elif ext == "ofx":
            texto = conteudo.decode("latin-1", errors="ignore")
            lancs = _ler_ofx(conteudo, nome)
        else:
            avisos.append(f"Formato '.{ext}' não suportado.")
    except Exception as e:  # arquivo corrompido, protegido por senha etc.
        avisos.append(f"Não foi possível ler o arquivo: {e}")

    if tabelas and not texto:
        texto = "\n".join(df.to_csv(sep=" ", index=False) for df in tabelas)

    tipo = tipo or detectar_tipo(nome, texto)
    doc = Documento(nome=nome, tipo=tipo, texto=texto, tabelas=tabelas, lancamentos=lancs, avisos=avisos)
    extrair_dados(doc, ano)
    return doc


def _totais_layout(doc: Documento):
    """Relatórios de período (anuais) com layout conhecido."""
    from .layouts import questor_icms_natureza

    return questor_icms_natureza(doc.texto, doc.nome)


def extrair_dados(doc: Documento, ano: int) -> None:
    """(Re)extrai lançamentos, séries e contas conforme o tipo do documento."""
    doc.series = {}
    doc.contas = {}
    doc.totais = []
    tipo = doc.tipo

    if tipo in (TipoDocumento.EXTRATO_CC, TipoDocumento.EXTRATO_APLICACAO, TipoDocumento.RENDIMENTOS):
        if not doc.lancamentos:
            for df in doc.tabelas:
                doc.lancamentos.extend(lancamentos_de_tabela(df, doc.nome, ano))
            if not doc.lancamentos and doc.texto:
                doc.lancamentos = lancamentos_de_texto(doc.texto, doc.nome, ano)
        if tipo == TipoDocumento.EXTRATO_CC:
            cred: dict[str, float] = defaultdict(float)
            deb: dict[str, float] = defaultdict(float)
            for l in doc.lancamentos:
                (cred if l.valor > 0 else deb)[competencia(l.data)] += abs(l.valor)
            doc.series = {"creditos_bancarios": dict(cred), "debitos_bancarios": dict(deb)}
            # Aplicação automática vinculada à conta corrente: rendimentos aparecem no próprio extrato.
            rend = _series_rendimentos(Documento(doc.nome, tipo, lancamentos=[
                l for l in doc.lancamentos if l.valor > 0 and "rend" in normalizar(l.descricao)
            ]), ano)
            doc.series.update(rend)
        else:
            doc.series = _series_rendimentos(doc, ano)
    elif tipo in (TipoDocumento.FATURAMENTO, TipoDocumento.COMPRAS, TipoDocumento.FISCAL, TipoDocumento.OUTRO) and (
        totais := _totais_layout(doc)
    ):
        doc.totais = totais
    elif tipo == TipoDocumento.FOLHA:
        doc.series = _series_folha(doc.texto, doc.tabelas, ano)
    elif tipo == TipoDocumento.FISCAL:
        doc.series = _series_fiscal(doc.texto, doc.tabelas, ano)
    elif tipo in (TipoDocumento.FATURAMENTO, TipoDocumento.COMPRAS):
        coluna = "faturamento" if tipo == TipoDocumento.FATURAMENTO else "compras"
        serie: dict[str, float] = {}
        for df in doc.tabelas:
            serie = serie_de_tabela(df, ano)
            if serie:
                break
        if not serie and doc.texto:
            serie = serie_de_texto(doc.texto, ano)
        doc.series = {coluna: serie} if serie else {}
    elif tipo == TipoDocumento.DEMONSTRATIVO:
        doc.contas = extrair_contas(doc.texto)

    if tipo != TipoDocumento.OUTRO and not doc.series and not doc.contas and not doc.lancamentos and not doc.totais:
        doc.avisos.append(
            "Nenhum valor foi reconhecido automaticamente. Confira o tipo do documento, use a leitura com IA "
            "ou digite os valores na tabela mensal."
        )
