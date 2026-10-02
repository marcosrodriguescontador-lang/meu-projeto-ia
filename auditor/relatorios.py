"""Geração dos relatórios (interno e para o cliente) em PDF com o timbre do escritório e em DOCX editável."""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from xml.sax.saxutils import escape

import pandas as pd
from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate, CondPageBreak, Frame, KeepTogether, NextPageTemplate, PageBreak, PageTemplate,
    Paragraph, Spacer, Table, TableStyle,
)

from .analises import Resultado
from .formatacao import brl, pct
from .modelos import Achado, Area, Documento, Empresa, Regime, Severidade

PASTA_ASSETS = Path(__file__).resolve().parent.parent / "assets"
TIMBRE = PASTA_ASSETS / "timbre.pdf"
ESCRITORIO = "Systema Serviços Contábeis e Financeiros"

VERMELHO = colors.HexColor("#91181B")
VERMELHO_CLARO = colors.HexColor("#F3E3DE")
CINZA = colors.HexColor("#4A4A4A")
CINZA_CLARO = colors.HexColor("#F4F4F4")
COR_SEVERIDADE = {
    Severidade.ALTA: colors.HexColor("#91181B"),
    Severidade.MEDIA: colors.HexColor("#C77700"),
    Severidade.BAIXA: colors.HexColor("#2F6EA5"),
    Severidade.INFO: colors.HexColor("#6B6B6B"),
}

# Áreas livres do timbre (em pontos, A4 = 595 x 842)
MARGEM_TOPO_CAPA = 4.6 * cm
MARGEM_TOPO = 4.0 * cm
MARGEM_BASE = 3.6 * cm
MARGEM_LADO = 2.0 * cm


@dataclass
class DadosRelatorio:
    empresa: Empresa
    resultado: Resultado
    documentos: list[Documento] = field(default_factory=list)
    parecer_interno: str = ""
    carta_cliente: str = ""
    analista: str = ""


# ---------------------------------------------------------------------------
# Fontes e estilos
# ---------------------------------------------------------------------------


def _registrar_fontes() -> tuple[str, str]:
    pasta = PASTA_ASSETS / "fonts"
    reg, bold = pasta / "Poppins-Regular.ttf", pasta / "Poppins-SemiBold.ttf"
    if not bold.exists():
        bold = pasta / "Poppins-Bold.ttf"
    if reg.exists() and bold.exists():
        try:
            pdfmetrics.registerFont(TTFont("Poppins", str(reg)))
            pdfmetrics.registerFont(TTFont("Poppins-Bold", str(bold)))
            return "Poppins", "Poppins-Bold"
        except Exception:
            pass
    return "Helvetica", "Helvetica-Bold"


FONTE, FONTE_BOLD = _registrar_fontes()


def _estilos() -> dict[str, ParagraphStyle]:
    base = ParagraphStyle("base", fontName=FONTE, fontSize=9.5, leading=13.5, textColor=CINZA, alignment=TA_JUSTIFY)
    return {
        "base": base,
        "pequeno": ParagraphStyle("pequeno", parent=base, fontSize=8, leading=10.5),
        "celula": ParagraphStyle("celula", parent=base, fontSize=8, leading=10, alignment=0),
        "celula_b": ParagraphStyle("celula_b", parent=base, fontSize=8, leading=10, alignment=0, fontName=FONTE_BOLD),
        "celula_dir": ParagraphStyle("celula_dir", parent=base, fontSize=8, leading=10, alignment=2),
        "celula_b_dir": ParagraphStyle("celula_b_dir", parent=base, fontSize=8, leading=10, alignment=2, fontName=FONTE_BOLD),
        "h1": ParagraphStyle("h1", parent=base, fontName=FONTE_BOLD, fontSize=15, leading=19, textColor=VERMELHO,
                             spaceBefore=6, spaceAfter=8, alignment=0),
        "h2": ParagraphStyle("h2", parent=base, fontName=FONTE_BOLD, fontSize=11.5, leading=15, textColor=VERMELHO,
                             spaceBefore=10, spaceAfter=5, alignment=0),
        "h3": ParagraphStyle("h3", parent=base, fontName=FONTE_BOLD, fontSize=10, leading=13, textColor=CINZA,
                             spaceBefore=4, spaceAfter=2, alignment=0),
        "capa_titulo": ParagraphStyle("capa_titulo", parent=base, fontName=FONTE_BOLD, fontSize=24, leading=29,
                                      textColor=VERMELHO, alignment=0),
        "capa_sub": ParagraphStyle("capa_sub", parent=base, fontSize=12, leading=17, alignment=0),
        "capa_tag": ParagraphStyle("capa_tag", parent=base, fontName=FONTE_BOLD, fontSize=9, leading=12,
                                   textColor=colors.white, alignment=0),
    }


E = _estilos()


def _p(texto: str, estilo: str = "base") -> Paragraph:
    return Paragraph(escape(texto).replace("\n", "<br/>"), E[estilo])


def _paragrafos(texto: str) -> list:
    """Converte texto livre (ex.: parecer da IA) em parágrafos; linhas curtas sem ponto final viram subtítulos."""
    saida = []
    for bloco in [b.strip() for b in texto.replace("\r", "").split("\n") if b.strip()]:
        if len(bloco) < 80 and not bloco.endswith((".", ":", ";")) and not bloco.startswith(("-", "•")):
            saida.append(_p(bloco, "h3"))
        else:
            saida.append(_p(bloco))
            saida.append(Spacer(1, 4))
    return saida


# ---------------------------------------------------------------------------
# Blocos comuns
# ---------------------------------------------------------------------------


def _tabela(dados: list[list], larguras: list[float], cabecalho: bool = True, alinhar_direita: list[int] | None = None) -> Table:
    direita = set(alinhar_direita or [])
    linhas = []
    for i, linha in enumerate(dados):
        estilo = "celula_b" if (cabecalho and i == 0) else "celula"
        linhas.append([
            c if not isinstance(c, str)
            else Paragraph(escape(c), E[estilo + "_dir"] if j in direita and not (cabecalho and i == 0) else E[estilo])
            for j, c in enumerate(linha)
        ])
    t = Table(linhas, colWidths=larguras, repeatRows=1 if cabecalho else 0)
    cmds = [
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#DDDDDD")),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    if cabecalho:
        cmds += [("BACKGROUND", (0, 0), (-1, 0), VERMELHO_CLARO), ("LINEBELOW", (0, 0), (-1, 0), 0.8, VERMELHO)]
    for i in range(1 if cabecalho else 0, len(linhas)):
        if i % 2 == 0:
            cmds.append(("BACKGROUND", (0, i), (-1, i), CINZA_CLARO))
    t.setStyle(TableStyle(cmds))
    return t


def _capa(titulo: str, subtitulo: str, d: DadosRelatorio, tag: str) -> list:
    emp = d.empresa
    regime = emp.regime.value + (f" - Anexo {emp.anexo_simples}" if emp.regime == Regime.SIMPLES else "")
    info = [
        ["Empresa", emp.nome or "-"],
        ["CNPJ", emp.cnpj or "-"],
        ["Regime tributário", regime],
        ["Atividade", emp.atividade.value],
        ["Período analisado", f"Ano-calendário {emp.ano_referencia}"],
        ["Data de emissão", date.today().strftime("%d/%m/%Y")],
    ]
    if d.analista:
        info.append(["Responsável", d.analista])
    tabela_info = Table([[_p(a, "celula_b"), _p(b, "celula")] for a, b in info], colWidths=[4.2 * cm, 9.5 * cm])
    tabela_info.setStyle(TableStyle([
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#DDDDDD")),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    etiqueta = Table([[_p(tag, "capa_tag")]], colWidths=[len(tag) * 0.2 * cm + 1 * cm])
    etiqueta.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), VERMELHO),
                                  ("LEFTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 3),
                                  ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
    etiqueta.hAlign = "LEFT"
    tabela_info.hAlign = "LEFT"
    return [
        Spacer(1, 1.2 * cm), etiqueta, Spacer(1, 0.5 * cm),
        _p(titulo, "capa_titulo"), Spacer(1, 0.25 * cm), _p(subtitulo, "capa_sub"),
        Spacer(1, 1.0 * cm), tabela_info,
        NextPageTemplate("normal"), PageBreak(),
    ]


def _bloco_achado(a: Achado, interno: bool) -> KeepTogether:
    cor = COR_SEVERIDADE[a.severidade]
    cab = Table(
        [[Paragraph(f"<font color='white'><b>{escape(a.severidade.value.upper())}</b></font>", E["celula"]),
          Paragraph(f"<b>{escape(a.titulo)}</b>", E["celula"]),
          Paragraph(escape(a.area.value), E["pequeno"])]],
        colWidths=[2.1 * cm, 10.6 * cm, 4.3 * cm],
    )
    cab.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, 0), cor), ("BACKGROUND", (1, 0), (-1, 0), CINZA_CLARO),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    corpo = [cab, Spacer(1, 3), _p(a.descricao)]
    corpo.append(Paragraph(f"<b>{'Recomendação' if interno else 'O que recomendamos'}:</b> {escape(a.recomendacao)}", E["base"]))
    if a.valor_envolvido:
        corpo.append(Paragraph(f"<b>Valor envolvido:</b> {brl(a.valor_envolvido)}", E["base"]))
    if a.fundamentacao:
        corpo.append(Paragraph(f"<i>Fundamentação: {escape(a.fundamentacao)}</i>", E["pequeno"]))
    if interno and not a.visivel_cliente:
        corpo.append(Paragraph("<i>Apontamento de uso interno (não consta no relatório do cliente).</i>", E["pequeno"]))
    corpo.append(Spacer(1, 8))
    return KeepTogether(corpo)


def _resumo_severidades(achados: list[Achado]) -> Table:
    contagem = [sum(1 for a in achados if a.severidade == s) for s in Severidade]
    celulas = []
    for s, n in zip(Severidade, contagem):
        celulas.append(Paragraph(
            f"<font size='16' color='{COR_SEVERIDADE[s].hexval().replace('0x', '#')}'><b>{n}</b></font><br/>"
            f"<font size='8'>{escape(s.value)}</font>", ParagraphStyle("c", parent=E["celula"], alignment=1, leading=18)))
    t = Table([celulas], colWidths=[4.25 * cm] * 4)
    t.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
                           ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
                           ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    return t


def _tabela_indicadores(ind: dict) -> Table:
    dados = [["Indicador", "Valor"]] + [[k, str(v)] for k, v in ind.items()]
    return _tabela(dados, [10 * cm, 7 * cm], alinhar_direita=[1])


def _tabela_simples(df: pd.DataFrame) -> Table:
    dados = [["Competência", "Receita", "RBT12", "Fator R", "Anexo", "Alíq. efetiva", "DAS calculado", "DAS pago"]]
    for _, r in df.iterrows():
        fr = r["Fator R"]
        dados.append([
            r["Competência"], brl(r["Receita do mês"]), brl(r["RBT12"]),
            pct(fr) if fr is not None and pd.notna(fr) else "-", r["Anexo aplicado"],
            pct(r["Alíquota efetiva"]), brl(r["DAS calculado"]),
            brl(r["DAS declarado/pago"]) if r["DAS declarado/pago"] else "-",
        ])
    dados.append(["Total", brl(df["Receita do mês"].sum()), "", "", "", "", brl(df["DAS calculado"].sum()),
                  brl(df["DAS declarado/pago"].sum()) if df["DAS declarado/pago"].sum() else "-"])
    return _tabela(dados, [1.9 * cm, 2.4 * cm, 2.5 * cm, 1.5 * cm, 1.3 * cm, 1.8 * cm, 2.4 * cm, 2.2 * cm],
                   alinhar_direita=[1, 2, 3, 5, 6, 7])


def _tabela_conciliacao(df: pd.DataFrame) -> Table:
    dados = [["Mês", "Faturamento", "Entradas bancos", "Não operac.", "Operacionais", "Diferença", "Op./Fat."]]
    for mes, r in df.iterrows():
        rel = r["Entradas / faturamento"]
        dados.append([mes, brl(r["Faturamento"]), brl(r["Entradas nos bancos"]), brl(r["Entradas não operacionais"]),
                      brl(r["Entradas operacionais"]), brl(r["Diferença"]), pct(rel, 0) if pd.notna(rel) else "-"])
    dados.append(["Total", brl(df["Faturamento"].sum()), brl(df["Entradas nos bancos"].sum()),
                  brl(df["Entradas não operacionais"].sum()), brl(df["Entradas operacionais"].sum()),
                  brl(df["Diferença"].sum()), ""])
    return _tabela(dados, [1.8 * cm, 2.6 * cm, 2.6 * cm, 2.4 * cm, 2.6 * cm, 2.5 * cm, 1.5 * cm],
                   alinhar_direita=[1, 2, 3, 4, 5, 6])


def _tabela_regimes(df: pd.DataFrame) -> Table:
    dados = [["Regime", "Tributos estimados (ano)", "% da receita", "Observação"]]
    for _, r in df.iterrows():
        dados.append([r["Regime"], brl(r["Tributos"]), pct(r["% da receita"]), r["Observação"]])
    return _tabela(dados, [4.4 * cm, 3.2 * cm, 2.0 * cm, 7.4 * cm], alinhar_direita=[1, 2])


# ---------------------------------------------------------------------------
# Montagem do PDF com timbre
# ---------------------------------------------------------------------------


def _numero_pagina(canvas, doc) -> None:
    if doc.page > 1:
        canvas.saveState()
        canvas.setFont(FONTE, 7.5)
        canvas.setFillColor(CINZA)
        canvas.drawRightString(A4[0] - MARGEM_LADO, MARGEM_BASE - 0.55 * cm, f"Página {doc.page}")
        canvas.restoreState()


def _montar_pdf(historia: list, titulo: str) -> bytes:
    buf = io.BytesIO()
    doc = BaseDocTemplate(buf, pagesize=A4, title=titulo, author=ESCRITORIO,
                          leftMargin=MARGEM_LADO, rightMargin=MARGEM_LADO, topMargin=MARGEM_TOPO, bottomMargin=MARGEM_BASE)
    largura = A4[0] - 2 * MARGEM_LADO
    capa = Frame(MARGEM_LADO, MARGEM_BASE, largura, A4[1] - MARGEM_TOPO_CAPA - MARGEM_BASE, id="capa")
    normal = Frame(MARGEM_LADO, MARGEM_BASE, largura, A4[1] - MARGEM_TOPO - MARGEM_BASE, id="normal")
    doc.addPageTemplates([PageTemplate("capa", [capa], onPage=_numero_pagina),
                          PageTemplate("normal", [normal], onPage=_numero_pagina)])
    doc.build(historia)
    return _aplicar_timbre(buf.getvalue())


def _aplicar_timbre(conteudo: bytes) -> bytes:
    from pypdf import PdfReader, PdfWriter

    if not TIMBRE.exists():
        return conteudo
    timbre_bytes = TIMBRE.read_bytes()
    paginas = PdfReader(io.BytesIO(conteudo)).pages
    escritor = PdfWriter()
    n_timbre = len(PdfReader(io.BytesIO(timbre_bytes)).pages)
    for i, pagina in enumerate(paginas):
        # Lê o timbre de novo a cada página para não compartilhar o mesmo objeto entre páginas.
        fundo = PdfReader(io.BytesIO(timbre_bytes)).pages[0 if i == 0 else min(1, n_timbre - 1)]
        escritor.add_page(fundo).merge_page(pagina)
    saida = io.BytesIO()
    escritor.write(saida)
    return saida.getvalue()


# ---------------------------------------------------------------------------
# Relatório interno
# ---------------------------------------------------------------------------


def relatorio_interno_pdf(d: DadosRelatorio) -> bytes:
    r = d.resultado
    h: list = _capa("Relatório Interno de Auditoria",
                    "Análise contábil, fiscal, tributária e de departamento pessoal", d, "USO INTERNO")

    h += [_p("1. Resumo executivo", "h1"), _resumo_severidades(r.achados), Spacer(1, 10)]
    if r.indicadores:
        h += [_p("Indicadores", "h2"), _tabela_indicadores(r.indicadores)]

    h += [CondPageBreak(6 * cm), _p("2. Apontamentos por área", "h1")]
    for area in Area:
        lista = [a for a in r.achados if a.area == area]
        if lista:
            h.append(_p(area.value, "h2"))
            h += [_bloco_achado(a, interno=True) for a in lista]
    if not r.achados:
        h.append(_p("Nenhum apontamento foi gerado com os dados disponíveis."))

    if d.parecer_interno:
        h += [PageBreak(), _p("3. Parecer técnico", "h1")] + _paragrafos(d.parecer_interno)

    h += [CondPageBreak(8 * cm), _p("4. Quadros de apoio", "h1")]
    if r.simples_mensal is not None and len(r.simples_mensal):
        h += [_p("Simples Nacional - apuração mensal recalculada", "h2"), _tabela_simples(r.simples_mensal)]
    if r.conciliacao_bancaria is not None and len(r.conciliacao_bancaria):
        h += [_p("Faturamento x movimentação bancária", "h2"), _tabela_conciliacao(r.conciliacao_bancaria)]
    if r.comparativo_regimes is not None:
        h += [_p("Comparativo de regimes tributários (simulação)", "h2"), _tabela_regimes(r.comparativo_regimes)]

    h += [CondPageBreak(6 * cm), _p("5. Documentos analisados", "h1")]
    dados_docs = [["Arquivo", "Tipo", "Observações"]]
    for doc in d.documentos:
        dados_docs.append([doc.nome, doc.tipo.value, " | ".join(doc.avisos) or "Lido sem ressalvas"])
    h.append(_tabela(dados_docs, [5 * cm, 5 * cm, 7 * cm]))
    if r.observacoes_metodo:
        h += [_p("Observações de método", "h2")] + [_p("- " + o) for o in r.observacoes_metodo]
    h += [Spacer(1, 10), _p(
        "Os cálculos são estimativas feitas a partir dos documentos fornecidos e da legislação parametrizada no "
        "sistema. Confirmar os pontos relevantes nos sistemas fiscais oficiais antes de qualquer retificação.", "pequeno")]
    return _montar_pdf(h, f"Relatório interno - {d.empresa.nome}")


# ---------------------------------------------------------------------------
# Relatório do cliente
# ---------------------------------------------------------------------------

TEXTO_PADRAO_CLIENTE = (
    "Apresentamos a seguir o resultado da análise das informações contábeis, fiscais, financeiras e de folha de "
    "pagamento da empresa referentes ao período indicado. O objetivo é identificar riscos, oportunidades de "
    "economia tributária e pontos de melhoria nos controles, de forma preventiva.\n"
    "Os pontos estão organizados por prioridade. Nossa equipe está à disposição para detalhar cada item e "
    "apoiar na implementação das recomendações."
)

INDICADORES_CLIENTE = [
    "Receita acumulada no ano", "Receita projetada para o ano", "RBT12 (último mês)", "Faixa atual",
    "Alíquota efetiva (último mês)", "Fator R (último mês)", "Anexo cadastrado", "DAS calculado no ano",
    "IRPJ estimado (ano)", "CSLL estimada (ano)", "PIS estimado (ano)", "COFINS estimada (ano)",
    "Rendimentos de aplicações (ano)",
]


def relatorio_cliente_pdf(d: DadosRelatorio) -> bytes:
    r = d.resultado
    visiveis = [a for a in r.achados if a.visivel_cliente]
    h: list = _capa("Relatório de Análise Fiscal e Tributária",
                    "Diagnóstico, pontos de atenção e recomendações", d, "RELATÓRIO AO CLIENTE")

    h += [_p("Apresentação", "h1")] + _paragrafos(d.carta_cliente or TEXTO_PADRAO_CLIENTE)

    ind = {k: r.indicadores[k] for k in INDICADORES_CLIENTE if k in r.indicadores}
    if ind:
        h += [_p("Números do período", "h2"), _tabela_indicadores(ind)]

    h += [CondPageBreak(6 * cm), _p("Pontos de atenção e recomendações", "h1")]
    if visiveis:
        h += [_bloco_achado(a, interno=False) for a in visiveis]
    else:
        h.append(_p("Não identificamos pontos de atenção relevantes com as informações analisadas."))

    if r.comparativo_regimes is not None:
        h += [CondPageBreak(6 * cm), _p("Comparativo de regimes tributários", "h2"),
              _p("Simulação com base no faturamento e na folha do período, sem considerar ICMS/ISS e benefícios "
                 "específicos. Serve como indicativo; a mudança de regime exige estudo detalhado.", "pequeno"),
              Spacer(1, 4), _tabela_regimes(r.comparativo_regimes)]
    if r.simples_mensal is not None and len(r.simples_mensal):
        h += [CondPageBreak(6 * cm), _p("Evolução da alíquota do Simples Nacional", "h2"), _tabela_simples(r.simples_mensal)]

    h += [Spacer(1, 14), _p(
        "Este relatório foi elaborado com base nos documentos e informações fornecidos pela empresa. Valores "
        "estimados podem variar conforme particularidades das operações. Em caso de dúvidas, fale com a nossa equipe.",
        "pequeno")]
    return _montar_pdf(h, f"Relatório de análise - {d.empresa.nome}")


# ---------------------------------------------------------------------------
# DOCX editável
# ---------------------------------------------------------------------------


def _imagens_timbre() -> tuple[bytes | None, bytes | None]:
    """Recorta cabeçalho e rodapé da página de continuação do timbre para usar no Word."""
    if not TIMBRE.exists():
        return None, None
    try:
        import pdfplumber

        with pdfplumber.open(str(TIMBRE)) as pdf:
            pagina = pdf.pages[min(1, len(pdf.pages) - 1)]
            img = pagina.to_image(resolution=200).original.convert("RGB")
        w, hgt = img.size
        cab, rod = img.crop((0, 0, w, int(hgt * 0.13))), img.crop((0, int(hgt * 0.905), w, hgt))
        saidas = []
        for parte in (cab, rod):
            b = io.BytesIO()
            parte.save(b, format="PNG")
            saidas.append(b.getvalue())
        return saidas[0], saidas[1]
    except Exception:
        return None, None


def relatorio_docx(d: DadosRelatorio, interno: bool) -> bytes:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    r = d.resultado
    documento = docx.Document()
    sec = documento.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin, sec.bottom_margin = Cm(4), Cm(3.6)
    sec.header_distance = sec.footer_distance = Cm(0)
    cab, rod = _imagens_timbre()
    if cab:
        p = sec.header.paragraphs[0]
        p.paragraph_format.left_indent = Cm(-2)
        p.add_run().add_picture(io.BytesIO(cab), width=Cm(21))
    if rod:
        p = sec.footer.paragraphs[0]
        p.paragraph_format.left_indent = Cm(-2)
        p.add_run().add_picture(io.BytesIO(rod), width=Cm(21))

    estilo = documento.styles["Normal"]
    estilo.font.name = "Poppins"
    estilo.font.size = Pt(10)
    vermelho = RGBColor(0x91, 0x18, 0x1B)

    def titulo(texto: str, nivel: int = 1):
        par = documento.add_heading(texto, level=nivel)
        for run in par.runs:
            run.font.color.rgb = vermelho
            run.font.name = "Poppins"

    titulo("Relatório Interno de Auditoria" if interno else "Relatório de Análise Fiscal e Tributária", 0)
    emp = d.empresa
    for rot, val in [("Empresa", emp.nome), ("CNPJ", emp.cnpj), ("Regime", emp.regime.value +
                     (f" - Anexo {emp.anexo_simples}" if emp.regime == Regime.SIMPLES else "")),
                     ("Período", f"Ano-calendário {emp.ano_referencia}"), ("Emissão", date.today().strftime("%d/%m/%Y"))]:
        par = documento.add_paragraph()
        par.add_run(f"{rot}: ").bold = True
        par.add_run(val or "-")

    if not interno:
        titulo("Apresentação")
        for bloco in (d.carta_cliente or TEXTO_PADRAO_CLIENTE).split("\n"):
            if bloco.strip():
                documento.add_paragraph(bloco.strip()).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    ind = r.indicadores if interno else {k: r.indicadores[k] for k in INDICADORES_CLIENTE if k in r.indicadores}
    if ind:
        titulo("Indicadores", 2)
        t = documento.add_table(rows=0, cols=2)
        t.style = "Light List Accent 2"
        for k, v in ind.items():
            c = t.add_row().cells
            c[0].text, c[1].text = k, str(v)

    titulo("Apontamentos" if interno else "Pontos de atenção e recomendações")
    for a in [x for x in r.achados if interno or x.visivel_cliente]:
        par = documento.add_paragraph()
        run = par.add_run(f"[{a.severidade.value.upper()}] {a.titulo}")
        run.bold = True
        run.font.color.rgb = vermelho if a.severidade == Severidade.ALTA else RGBColor(0x4A, 0x4A, 0x4A)
        documento.add_paragraph(a.descricao).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        par = documento.add_paragraph()
        par.add_run("Recomendação: ").bold = True
        par.add_run(a.recomendacao)
        if a.valor_envolvido:
            par = documento.add_paragraph()
            par.add_run("Valor envolvido: ").bold = True
            par.add_run(brl(a.valor_envolvido))
        if a.fundamentacao:
            documento.add_paragraph(f"Fundamentação: {a.fundamentacao}").runs[0].italic = True

    if interno and d.parecer_interno:
        titulo("Parecer técnico")
        for bloco in d.parecer_interno.split("\n"):
            if bloco.strip():
                documento.add_paragraph(bloco.strip()).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if r.comparativo_regimes is not None:
        titulo("Comparativo de regimes tributários (simulação)", 2)
        t = documento.add_table(rows=1, cols=3)
        t.style = "Light List Accent 2"
        for c, txt in zip(t.rows[0].cells, ["Regime", "Tributos (ano)", "% da receita"]):
            c.text = txt
        for _, linha in r.comparativo_regimes.iterrows():
            c = t.add_row().cells
            c[0].text, c[1].text, c[2].text = linha["Regime"], brl(linha["Tributos"]), pct(linha["% da receita"])

    buf = io.BytesIO()
    documento.save(buf)
    return buf.getvalue()
